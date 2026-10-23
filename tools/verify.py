"""One command that checks the whole deliverable is self-consistent.

The acceptance list in ``docs/acceptance.md`` is generated from what this script
verifies, so a reviewer can reproduce every tick instead of trusting the report.
It is deliberately cheap: no GPU, no checkpoint, no trained detector. The heavy
paths (metrics, latency) are verified by reading the committed result files and
comparing them against the published paper numbers.

Example:
    python tools/verify.py            # human-readable checklist
    python tools/verify.py --json     # machine-readable, for CI
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import subprocess
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
FAILS: list[str] = []


def check(name: str, ok: bool, detail: str = "") -> bool:
    print(f"[{'PASS' if ok else 'FAIL'}] {name}" + (f"  — {detail}" if detail else ""))
    if not ok:
        FAILS.append(name)
    return ok


def read_csv(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with path.open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


def check_paper_metrics() -> None:
    """The four headline numbers must equal the paper / README exactly."""
    expected = {"S": 0.8718, "MAE": 0.0501, "maxE": 0.9475, "maxF": 0.9129}
    rows = {r["Model"]: r for r in read_csv(ROOT / "predictions" / "metrics.csv")}
    ours = rows.get("Ours")
    if not check("predictions/metrics.csv 含 Ours 行", ours is not None):
        return
    for key, value in expected.items():
        got = float(ours[key])
        check(f"复现指标 {key} = {value}", abs(got - value) < 5e-5, f"读到 {got}")
    check("对比方法共 12 行", len(rows) == 12, f"读到 {len(rows)} 行")

    abl = read_csv(ROOT / "ablation" / "metrics.csv")
    check("消融共 4 个变体", len(abl) == 4, f"读到 {len(abl)} 行")
    if abl:
        worst = min(abl, key=lambda r: float(r["S"]))
        check("消融中最差变体是去掉 Transformer",
              worst["Model"] == "w_o_Transformer", f"读到 {worst['Model']}")


def check_system_eval() -> None:
    path = ROOT / "docs" / "assets" / "system-eval.json"
    if not check("docs/assets/system-eval.json 存在", path.exists()):
        return
    payload = json.loads(path.read_text(encoding="utf-8"))
    mask = payload.get("mask", {})
    lat = payload.get("latency_ms", {})
    check("系统评测覆盖全量 362 张", payload.get("images") == 362,
          f"读到 {payload.get('images')}")
    check("系统评测无失败", payload.get("errors") == 0, f"读到 {payload.get('errors')}")
    check("掩膜 IoU 均值 > 0.8", float(mask.get("iou_mean", 0)) > 0.8,
          f"读到 {mask.get('iou_mean')}")
    check("端到端 P50 < 300 ms", float(lat.get("total_p50", 1e9)) < 300,
          f"读到 {lat.get('total_p50')}")
    verdicts = payload.get("verdicts", {})
    check("判定三档均有样本",
          all(verdicts.get(k, 0) > 0 for k in ("light", "medium", "severe")),
          json.dumps(verdicts, ensure_ascii=False))


def check_deploy() -> None:
    path = ROOT / "docs" / "assets" / "latency.json"
    if not check("docs/assets/latency.json 存在", path.exists()):
        return
    payload = json.loads(path.read_text(encoding="utf-8"))
    rows = payload.get("latency", [])
    check("延迟表含 4 个后端", len(rows) == 4, f"读到 {len(rows)}")
    fid = payload.get("fidelity", {})
    fp16 = fid.get("torch-fp16", {})
    check("fp16 相对 fp32 的 mask IoU ≥ 0.95",
          float(fp16.get("mask_iou_vs_fp32_min", 0)) >= 0.95,
          f"读到 {fp16.get('mask_iou_vs_fp32_min')}")
    check("ONNX fp32 与 torch fp32 一致（IoU > 0.999）",
          float(fid.get("onnx-fp32", {}).get("mask_iou_vs_fp32_mean", 0)) > 0.999,
          str(fid.get("onnx-fp32", {}).get("mask_iou_vs_fp32_mean")))


def check_pr() -> None:
    rows = read_csv(ROOT / "docs" / "assets" / "pr-metrics.csv")
    if not check("docs/assets/pr-metrics.csv 存在", bool(rows)):
        return
    ours = next((r for r in rows if r["method"] == "Ours"), None)
    if check("PR 表含 Ours", ours is not None):
        check("PR 的 maxF 与 metrics.csv 一致", abs(float(ours["maxF"]) - 0.9129) < 1e-3,
              f"读到 {ours['maxF']}")


def check_docs() -> None:
    docs = sorted((ROOT / "docs").glob("*.md"))
    check("docs 至少 10 篇", len(docs) >= 10, f"读到 {len(docs)} 篇")
    index = (ROOT / "docs" / "README.md").read_text(encoding="utf-8")
    missing = [p.name for p in docs
               if p.name != "README.md" and p.name not in index]
    check("索引链接到全部文档", not missing, "缺少：" + ", ".join(missing) if missing else "")

    assets = ROOT / "docs" / "assets"
    for name in ("system-architecture.drawio", "glide-net-architecture.drawio",
                 "inspection-pipeline.drawio"):
        path = assets / name
        if not check(f"架构源文件 {name} 存在", path.exists()):
            continue
        try:
            ET.parse(path)
            check(f"{name} 是合法 XML", True)
        except ET.ParseError as exc:
            # A malformed file still exports a PNG — just an almost empty one.
            check(f"{name} 是合法 XML", False, str(exc))
    broken = []
    for md in docs:
        text = md.read_text(encoding="utf-8")
        for target in re.findall(r"\]\((?!http)([^)#\s]+)", text):
            if not (md.parent / target).exists():
                broken.append(f"{md.name} -> {target}")
    check("文档内部链接无 404", not broken, "; ".join(broken[:4]))


def check_gitignore() -> None:
    ignored = ["ckpt/glide_net_6440647.pth", "materials/Manuscript.zip",
               "runs/detect/x/y.pt", "deploy/onnx/glide_net_352_fp16.onnx",
               "results/batch_100/batch.csv"]
    cmd = ["git", "check-ignore", "--no-index", *ignored]
    proc = subprocess.run(cmd, cwd=ROOT, capture_output=True, text=True)
    hit = {line.strip() for line in proc.stdout.splitlines() if line.strip()}
    check("大文件与课程材料全部被忽略", len(hit) == len(ignored),
          f"{len(hit)}/{len(ignored)}")
    if hit != set(ignored):
        check("未被忽略的路径", False, ", ".join(sorted(set(ignored) - hit)))

    log = subprocess.run(["git", "log", "--all", "--pretty=format:", "--name-only"],
                         cwd=ROOT, capture_output=True, text=True).stdout
    leaked = [line for line in log.splitlines()
              if line.startswith(("ckpt/", "materials/"))]
    check("git 历史中不含 ckpt/ 与 materials/", not leaked,
          f"发现 {len(leaked)} 个文件")


def check_tests() -> None:
    proc = subprocess.run([sys.executable, "-m", "pytest", "tests", "-q"],
                          cwd=ROOT, capture_output=True, text=True)
    tail = (proc.stdout + proc.stderr).strip().splitlines()[-1] if proc.returncode != -1 else ""
    check("单元测试全绿", proc.returncode == 0, tail)
    m = re.search(r"(\d+) passed", tail)
    if m:
        check("测试数量 ≥ 20", int(m.group(1)) >= 20, tail)


def check_assets() -> None:
    assets = ROOT / "docs" / "assets"
    required = ["system-architecture.png", "glide-net-architecture.png",
                "inspection-pipeline.png", "training-curves.png",
                "comparison-bars.png", "ablation-bars.png", "pr-curves.png",
                "latency.png", "severity-dist.png", "gallery.png",
                "example-panel.png", "example-overlay.png", "ui-workstation.png"]
    missing = [n for n in required if not (assets / n).exists()]
    check("报告所需图表齐全", not missing, "缺少：" + ", ".join(missing) if missing else "")


def check_slides() -> None:
    deck = ROOT / "docs" / "slides.pptx"
    if not check("docs/slides.pptx 存在", deck.exists()):
        return
    try:
        from pptx import Presentation
    except ImportError:
        check("python-pptx 可导入", False, "pip install python-pptx")
        return
    try:
        prs = Presentation(deck)
        n = len(prs.slides._sldIdLst)
    except Exception as exc:  # a corrupt package must fail loudly
        check("slides.pptx 可解析", False, str(exc))
        return
    check("slides.pptx 可解析", True, f"{n} 页")
    check("幻灯片页数为 14", n == 14, f"读到 {n} 页")
    empty = [i for i, sl in enumerate(prs.slides, 1)
             if not "".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame).strip()]
    check("每页都有文字", not empty, "空白页：" + str(empty) if empty else "")

    # the deck is generated, so a stale build is the failure mode we care about
    try:
        sys.path.insert(0, str(ROOT / "tools"))
        import check_slides
        problems = check_slides.check(deck)
    except Exception as exc:
        check("幻灯片版式自检可运行", False, str(exc))
        return
    check("幻灯片无溢出 / 重叠 / 越界", not problems, "; ".join(problems[:3]))


CHECKS = [check_paper_metrics, check_system_eval, check_deploy, check_pr,
          check_docs, check_gitignore, check_assets, check_slides, check_tests]


def main() -> int:
    ap = argparse.ArgumentParser(description="Verify the course-practice deliverable")
    ap.add_argument("--json", action="store_true", help="print a JSON summary")
    args = ap.parse_args()
    print(f"verify: {ROOT}\n")
    for fn in CHECKS:
        print(f"--- {fn.__name__}")
        fn()
        print()
    total = len(FAILS)
    if args.json:
        print(json.dumps({"failed": FAILS, "passed": total == 0}, ensure_ascii=False))
    else:
        print("=" * 60)
        print("全部通过 ✅" if total == 0 else f"{total} 项未通过 ❌: {', '.join(FAILS)}")
    return 0 if total == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
